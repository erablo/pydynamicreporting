"""
Generate a PowerPoint overview of PMD (Packaging and Marking Development) at Ericsson:
responsibilities, people and projects.

Styling follows the Ericsson / Templafy "Atlas" template conventions (title slide,
"Ericsson Internal" security class, footer with signum / revision / date, Ericsson
palette). Run Templafy's "Apply / Refresh template" inside PowerPoint afterwards to
bind the official branded master if required.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------- brand palette
ERICSSON_BLUE = RGBColor(0x00, 0x30, 0x56)
ERICSSON_ACCENT = RGBColor(0x00, 0x82, 0xF0)
ERICSSON_GRAY = RGBColor(0x53, 0x53, 0x5A)
LIGHT_GRAY = RGBColor(0xF2, 0xF2, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL = RGBColor(0x00, 0xA9, 0xA5)
PURPLE = RGBColor(0x62, 0x2D, 0x6E)

FONT = "Arial"
PREPARED = "ERABLO Anders Blom"
REVISION = "PA1"
DATE = "2026-10-12"
SECURITY = "Ericsson Internal"
TOTAL_PAGES = 9

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


# ------------------------------------------------------------------- helpers
def add_rect(slide, x, y, w, h, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_text(slide, x, y, w, h, text, size, color=ERICSSON_GRAY, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT, italic=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color
    return tb


def add_footer(slide, page_no):
    add_rect(slide, Inches(0.5), Inches(7.02), Inches(12.33), Pt(1), ERICSSON_ACCENT)
    footer = f"{PREPARED}  |  {REVISION}  |  {DATE}  |  {SECURITY}  |  Page {page_no} of {TOTAL_PAGES}"
    add_text(slide, Inches(0.5), Inches(7.08), Inches(12.33), Inches(0.3),
             footer, 8, ERICSSON_GRAY, align=PP_ALIGN.LEFT)


def content_header(slide, title, page_no):
    add_rect(slide, Inches(0.5), Inches(0.35), Inches(0.12), Inches(0.55), ERICSSON_ACCENT)
    add_text(slide, Inches(0.75), Inches(0.3), Inches(11.9), Inches(0.7),
             title, 26, ERICSSON_BLUE, bold=True)
    add_rect(slide, Inches(0.5), Inches(1.05), Inches(12.33), Pt(1.5), LIGHT_GRAY)
    add_footer(slide, page_no)


def bullets(slide, x, y, w, h, items, size=14, gap=6):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (txt, lvl) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = lvl
        p.space_after = Pt(gap)
        marker = "\u2022  " if lvl == 0 else "\u2013  "
        r = p.add_run()
        r.text = marker + txt
        r.font.size = Pt(size - lvl * 2)
        r.font.name = FONT
        r.font.color.rgb = ERICSSON_BLUE if lvl == 0 else ERICSSON_GRAY
        r.font.bold = (lvl == 0)
    return tb


def add_table(slide, x, y, w, headers, rows, col_widths,
              body_size=11, header_size=11, row_h=0.4, header_color=ERICSSON_BLUE):
    n_rows = len(rows) + 1
    h = Inches(row_h * n_rows)
    tbl = slide.shapes.add_table(n_rows, len(headers), x, y, w, h).table
    for ci, cw in enumerate(col_widths):
        tbl.columns[ci].width = cw
    for ci, head in enumerate(headers):
        cell = tbl.cell(0, ci)
        cell.fill.solid(); cell.fill.fore_color.rgb = header_color
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        r = p.add_run(); r.text = head
        r.font.size = Pt(header_size); r.font.bold = True
        r.font.color.rgb = WHITE; r.font.name = FONT
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if ri % 2 else LIGHT_GRAY
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text_frame.word_wrap = True
            p = cell.text_frame.paragraphs[0]
            r = p.add_run(); r.text = str(val)
            r.font.size = Pt(body_size)
            r.font.color.rgb = ERICSSON_GRAY; r.font.name = FONT
            r.font.bold = (ci == 0)
    return tbl


def card(slide, x, y, w, h, title, body, accent):
    add_rect(slide, x, y, w, h, LIGHT_GRAY)
    add_rect(slide, x, y, Inches(0.1), h, accent)
    add_text(slide, Emu(int(x) + Inches(0.3)), Emu(int(y) + Inches(0.15)),
             Emu(int(w) - Inches(0.5)), Inches(0.5), title, 14, ERICSSON_BLUE, bold=True)
    add_text(slide, Emu(int(x) + Inches(0.3)), Emu(int(y) + Inches(0.7)),
             Emu(int(w) - Inches(0.5)), Emu(int(h) - Inches(0.85)), body, 11, ERICSSON_GRAY)


# ============================================================= SLIDE 1 - TITLE
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, ERICSSON_BLUE)
add_rect(s, 0, Inches(4.9), SW, Inches(0.08), ERICSSON_ACCENT)
add_rect(s, Inches(0.9), Inches(2.0), Inches(0.18), Inches(2.4), TEAL)
add_text(s, Inches(1.25), Inches(2.0), Inches(11.0), Inches(1.3),
         "Packaging & Marking Development", 40, WHITE, bold=True)
add_text(s, Inches(1.25), Inches(3.15), Inches(11.0), Inches(0.8),
         "PMD \u2013 Organization Overview", 24, RGBColor(0xBF, 0xD9, 0xF2))
add_text(s, Inches(1.25), Inches(5.15), Inches(11.0), Inches(1.0),
         "Responsibilities  |  People  |  Projects\n"
         "BNEW RTE / EU Hardware  \u2013  Head: Sanna Severinsson", 15, WHITE)
add_text(s, Inches(0.9), Inches(6.9), Inches(11.5), Inches(0.4),
         f"{SECURITY}   |   {REVISION}   |   {DATE}   |   Prepared by {PREPARED}",
         10, RGBColor(0x9F, 0xC2, 0xE8))
add_text(s, Inches(11.0), Inches(0.4), Inches(1.9), Inches(0.4),
         "ERICSSON", 16, WHITE, bold=True, align=PP_ALIGN.RIGHT)

# ================================================= SLIDE 2 - WHAT IS PMD
s = prs.slides.add_slide(BLANK)
content_header(s, "What is PMD?", 2)
add_text(s, Inches(0.75), Inches(1.4), Inches(11.8), Inches(1.1),
         "PMD = Packaging and Marking Development \u2013 a specialist organization within "
         "BNEW RTE / EU Hardware. It is Ericsson's engineering and governance center for "
         "packaging materials and the global owner of product marking and labelling "
         "standards.", 15, ERICSSON_GRAY)
card(s, Inches(0.75), Inches(2.8), Inches(3.75), Inches(3.6), "Engineering center",
     "Designs and industrializes protective packaging across the full material "
     "lifecycle for Ericsson and supplier products.", TEAL)
card(s, Inches(4.8), Inches(2.8), Inches(3.75), Inches(3.6), "Governance owner",
     "Corporate category owner for product marking & labelling; sets global "
     "standards and approves exceptions.", ERICSSON_ACCENT)
card(s, Inches(8.85), Inches(2.8), Inches(3.75), Inches(3.6), "Competence hub",
     "Drives packaging requirements, design methods, compliance, simulation "
     "capability and continuous improvement.", PURPLE)

# ================================================= SLIDE 3 - RESPONSIBILITIES
s = prs.slides.add_slide(BLANK)
content_header(s, "Responsibilities & Scope", 3)
add_text(s, Inches(0.75), Inches(1.25), Inches(5.6), Inches(0.4),
         "Packaging development (full lifecycle)", 15, ERICSSON_BLUE, bold=True)
bullets(s, Inches(0.75), Inches(1.75), Inches(5.7), Inches(5.0), [
    ("Concepts, material selection, mechanical design, prototyping, verification, "
     "release & maintenance", 0),
    ("White-box (Ericsson) and black-box (supplier) products, inbound materials, "
     "transport packaging, fitments, pallets, bags", 0),
    ("Requirements for handling, stacking, transport, product protection, ESD, "
     "moisture, dangerous goods, recycling", 0),
    ("Packaging-design review & approval", 0),
    ("Coordination with product design, PE, supply, sourcing, PM, suppliers & "
     "manufacturing sites", 0),
], size=12.5, gap=7)
add_text(s, Inches(6.7), Inches(1.25), Inches(6.0), Inches(0.4),
         "Marking & labelling (category owner)", 15, ERICSSON_BLUE, bold=True)
bullets(s, Inches(6.7), Inches(1.75), Inches(6.0), Inches(5.0), [
    ("Product, package, handling-unit, transport, conformity, recycling, "
     "supplier-ID and barcode / 2D-code markings", 0),
    ("Conformity marks: CE, UKCA, FCC and market-specific", 0),
    ("Standards for placement, free zones, readability & barcode quality "
     "(ISO 11469, ISO 7000, ISPM 15)", 0),
    ("Approves exceptions to standard packaging markings", 0),
    ("Governed via cross-functional steering group (GFTL, PA, R&D, Supply, "
     "Traceability, Sourcing, Environmental)", 0),
], size=12.5, gap=7)

# ================================================= SLIDE 4 - LEADERSHIP/ORG
s = prs.slides.add_slide(BLANK)
content_header(s, "People \u2013 Leadership & Team", 4)
add_rect(s, Inches(4.9), Inches(1.35), Inches(3.55), Inches(0.75), ERICSSON_BLUE)
add_text(s, Inches(4.9), Inches(1.35), Inches(3.55), Inches(0.75),
         "Sanna Severinsson\nHead of Packaging & Marking Development", 12, WHITE,
         bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_rect(s, Inches(6.6), Inches(2.1), Pt(1.5), Inches(0.3), ERICSSON_ACCENT)
add_text(s, Inches(0.75), Inches(2.55), Inches(11.8), Inches(0.35),
         "12 direct reports (Sweden), plus contributors across China & supply sites",
         12, ERICSSON_GRAY, italic=True, align=PP_ALIGN.CENTER)
rows = [
    ("Anders Blom", "Product Development Leader \u2013 packaging simulation"),
    ("Anna-Stina Nilsson", "Technical Area Leader, Packaging Material"),
    ("Ismail Sebugwawo", "Senior Developer, Packaging Material"),
    ("Marija Kusic", "Senior Developer, Packaging Material"),
    ("Lucia Alhalabi", "Packaging Engineer \u2013 molded pulp & simulation"),
    ("Kenneth Wing\u00e5rdh", "Product Development Leader \u2013 HPS / tensile testing"),
    ("Magnus Lidevi", "Master Developer, Marking & Labelling"),
    ("Jessica Guo", "Product Development Leader, M&L"),
    ("Magnus Lindberg Q", "Packaging Cost & Quality Engineer"),
    ("Marcin Czechowski", "Technical Product Manager"),
    ("Robert Lozano", "Operational Product Owner"),
    ("Tony Westlund", "Maintenance & Quality Driver, Packaging"),
]
add_table(s, Inches(0.75), Inches(3.0), Inches(11.8),
          ["Name", "Role"], rows,
          [Inches(3.3), Inches(8.5)], body_size=10.5, header_size=11, row_h=0.31)

# ================================================= SLIDE 5 - PROJECTS OVERVIEW
s = prs.slides.add_slide(BLANK)
content_header(s, "Projects & Initiatives", 5)
card(s, Inches(0.75), Inches(1.4), Inches(5.85), Inches(2.5),
     "Packaging Simulation Pilot",
     "Ansys / LS-Dyna drop simulation for fiber & molded-pulp packaging. Virtual-lab "
     "material cards (MAT_PAPER \u2192 MAT_058), 6-drop templates, 19-inch baseband "
     "demonstrator. ~8.67 MSEK annual value; commercial licensing approved Dec 2025.",
     TEAL)
card(s, Inches(6.75), Inches(1.4), Inches(5.85), Inches(2.5),
     "Molded-pulp & fiber packaging",
     "Flagship sustainability program replacing plastic fitments with recycled fiber, "
     "corrugated board, molded pulp and honeycomb across radios and enclosures.",
     ERICSSON_ACCENT)
card(s, Inches(0.75), Inches(4.1), Inches(5.85), Inches(2.5),
     "Marking & labelling maintenance",
     "Maintains standard label solutions (e.g. SVF 191 050 variants), 2D-code "
     "adoption, conformity marks and the global marking framework.",
     PURPLE)
card(s, Inches(6.75), Inches(4.1), Inches(5.85), Inches(2.5),
     "Supplier collaboration",
     "Concept-to-release work with Nefab, DS Smith, Ansys, ESS Nanjing and IKEA on "
     "material data, testing, simulation and industrialization.",
     ERICSSON_GRAY)

# ================================================= SLIDE 6 - MOLDED PULP
s = prs.slides.add_slide(BLANK)
content_header(s, "Molded-Pulp & Fiber Packaging \u2013 Examples", 6)
rows = [
    ("AIR3219 B77U", "Corrugated + molded pulp; launched China Mar 2026; plastic down "
                     "to 1.83%; reused tooling & auto-packing lines"),
    ("Enclosure 6180 Hex\n(RTK 294 311/1)", "PMD's first plastic-free black-box "
                                            "enclosure (corrugated + honeycomb)"),
    ("AIR3255 / AIR3265", "Molded-pulp fitment solutions through verification & release"),
    ("Stockholm Up 6628", "Molded-pulp package tested via drop, vibration, compression "
                          "& stacking \u2013 no product damage reported"),
    ("40 kg radio packaging", "100% recycled molded pulp demonstrated (ESS Nanjing)"),
]
add_table(s, Inches(0.75), Inches(1.5), Inches(11.8),
          ["Product", "Solution & status"], rows,
          [Inches(3.3), Inches(8.5)], body_size=11.5, header_size=12, row_h=0.78)
add_text(s, Inches(0.75), Inches(6.4), Inches(11.8), Inches(0.5),
         "Note: solutions are not automatically plastic-free \u2013 some still use PE "
         "abrasion film, tape or small plastic components; PMD works to reduce residual "
         "plastic while retaining product protection.", 11, ERICSSON_GRAY, italic=True)

# ================================================= SLIDE 7 - SUSTAINABILITY
s = prs.slides.add_slide(BLANK)
content_header(s, "Sustainability Goals", 7)
add_text(s, Inches(0.75), Inches(1.35), Inches(11.8), Inches(0.8),
         "Core objective: substantially cut plastic content, replacing it with "
         "recyclable wood-fiber solutions.", 15, ERICSSON_BLUE, bold=True)
# staged target bar
targets = [("2026", "10%"), ("2027", "5%"), ("2028", "1%"), ("2030", "<1\u20132%")]
bw = Inches(2.6); y = Inches(2.4); startx = Inches(1.1); gapx = Inches(0.4)
cols = [TEAL, ERICSSON_ACCENT, ERICSSON_BLUE, PURPLE]
for i, (yr, pct) in enumerate(targets):
    bx = Emu(int(startx) + i * (int(bw) + int(gapx)))
    add_rect(s, bx, y, bw, Inches(1.3), cols[i])
    add_text(s, bx, Emu(int(y) + Inches(0.15)), bw, Inches(0.6), pct, 26, WHITE,
             bold=True, align=PP_ALIGN.CENTER)
    add_text(s, bx, Emu(int(y) + Inches(0.8)), bw, Inches(0.4), "plastic by " + yr, 12,
             WHITE, align=PP_ALIGN.CENTER)
add_text(s, Inches(0.75), Inches(4.0), Inches(11.8), Inches(0.5),
         "Targets vary across requirement versions (staged 10/5/1% vs. <1% or <2% by "
         "2030) \u2013 check the applicable project requirement.", 11, ERICSSON_GRAY,
         italic=True)
bullets(s, Inches(0.75), Inches(4.6), Inches(11.8), Inches(2.2), [
    ("Compliance with EU Packaging Directive 94/62/EC and ISO 18601", 0),
    ("Ericsson banned/restricted-substance requirements; no PVC or halogenated flame "
     "retardants; supplier material-content declarations", 0),
    ("Recycling & material-identification marking; transport-fill optimization; reuse, "
     "salvage and plastic-free designs where feasible", 0),
], size=13, gap=8)

# ================================================= SLIDE 8 - SUPPLIERS
s = prs.slides.add_slide(BLANK)
content_header(s, "Supplier Collaborations", 8)
rows = [
    ("Nefab", "Molded-pulp / fiber development, material & geometry data, transport "
              "testing, simulation (Nefab China Test Lab)"),
    ("DS Smith", "Paper & corrugated samples, grammage / thickness data, tensile-test "
                 "inputs for fiber simulation"),
    ("Ansys", "LS-Dyna simulation-tool support, templates, automation & pilot methodology"),
    ("ESS Nanjing", "Industrialization of recycled molded-pulp packaging; high-weight "
                    "radio protection"),
    ("IKEA & others", "Experience sharing on molded-pulp simulation & material modelling"),
]
add_table(s, Inches(0.75), Inches(1.5), Inches(11.8),
          ["Partner", "Collaboration"], rows,
          [Inches(2.8), Inches(9.0)], body_size=12, header_size=12, row_h=0.72)
add_text(s, Inches(0.75), Inches(6.1), Inches(11.8), Inches(0.6),
         "Supplier-facing process: issue a Packaging Design Assignment \u2192 review "
         "proposals \u2192 select partner with Sourcing \u2192 approve prototypes \u2192 "
         "obtain material declarations.", 12, ERICSSON_BLUE, bold=True)

# ================================================= SLIDE 9 - CLOSING
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, ERICSSON_BLUE)
add_rect(s, Inches(0.9), Inches(2.5), Inches(0.18), Inches(1.9), TEAL)
add_text(s, Inches(1.25), Inches(2.6), Inches(10.8), Inches(1.7),
         "PMD: designing protective, sustainable\npackaging and governing Ericsson's\n"
         "global marking & labelling", 26, WHITE, bold=True)
add_text(s, Inches(1.25), Inches(4.6), Inches(10.8), Inches(0.6),
         "Thank you", 20, RGBColor(0xBF, 0xD9, 0xF2))
add_text(s, Inches(0.9), Inches(6.9), Inches(11.5), Inches(0.4),
         f"{SECURITY}   |   {REVISION}   |   {DATE}   |   {PREPARED}",
         10, RGBColor(0x9F, 0xC2, 0xE8))

# -------------------------------------------------------------------- save
import sys
out = sys.argv[1] if len(sys.argv) > 1 else "PMD_Overview.pptx"
prs.save(out)
print("Saved:", out)
print("Slides:", len(prs.slides._sldIdLst))
